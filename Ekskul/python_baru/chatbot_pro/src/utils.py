import base64
import pypdf
from io import BytesIO

def encode_image(uploaded_file):
    """Encodes an uploaded image file to a base64 string."""
    return base64.b64encode(uploaded_file.getvalue()).decode('utf-8')

def extract_text_from_pdf(uploaded_file):
    """Extracts text content from an uploaded PDF file."""
    try:
        pdf_reader = pypdf.PdfReader(uploaded_file)
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text() + "\n"
        return text
    except Exception as e:
        return f"Error membaca PDF: {e}"

def extract_text_from_docx(uploaded_file):
    """Extracts text content from a DOCX file."""
    try:
        from docx import Document
        doc = Document(uploaded_file)
        text = "\n".join([para.text for para in doc.paragraphs])
        return text
    except ImportError:
        return "[Error: python-docx not installed. Run: pip install python-docx]"
    except Exception as e:
        return f"Error membaca DOCX: {e}"

def extract_text_from_csv(uploaded_file):
    """Extracts text content from a CSV file."""
    try:
        import csv
        from io import StringIO
        content = uploaded_file.getvalue().decode("utf-8")
        reader = csv.reader(StringIO(content))
        rows = list(reader)
        # Format as readable text
        text = "\n".join([", ".join(row) for row in rows])
        return text
    except Exception as e:
        return f"Error membaca CSV: {e}"

def extract_text_from_xlsx(uploaded_file):
    """Extracts text content from an XLSX file."""
    try:
        import openpyxl
        from io import BytesIO
        wb = openpyxl.load_workbook(BytesIO(uploaded_file.getvalue()), data_only=True)
        text = ""
        for sheet in wb.worksheets:
            text += f"\n[Sheet: {sheet.title}]\n"
            for row in sheet.iter_rows(values_only=True):
                row_text = ", ".join([str(cell) if cell else "" for cell in row])
                text += row_text + "\n"
        return text
    except ImportError:
        return "[Error: openpyxl not installed. Run: pip install openpyxl]"
    except Exception as e:
        return f"Error membaca XLSX: {e}"

def extract_text_from_html(uploaded_file):
    """Extracts text content from an HTML file."""
    try:
        from html.parser import HTMLParser
        class TextExtractor(HTMLParser):
            def __init__(self):
                super().__init__()
                self.text = []
            def handle_data(self, data):
                self.text.append(data.strip())
        
        content = uploaded_file.getvalue().decode("utf-8")
        parser = TextExtractor()
        parser.feed(content)
        return " ".join(filter(None, parser.text))
    except Exception as e:
        return f"Error membaca HTML: {e}"

def extract_text_from_json(uploaded_file):
    """Extracts text content from a JSON file."""
    try:
        import json
        content = uploaded_file.getvalue().decode("utf-8")
        data = json.loads(content)
        # Pretty print JSON as text
        return json.dumps(data, indent=2, ensure_ascii=False)
    except Exception as e:
        return f"Error membaca JSON: {e}"

def get_smart_pdf_context(full_text, user_query, max_chars=10000):
    """
    Simulates a lightweight RAG by taking the first 2000 chars 
    AND searching for relevant paragraphs that contain query keywords.
    """
    # 1. Always include the first section (introduction/abstract)
    intro_size = 2000
    context = full_text[:intro_size]
    
    # 2. Simple Keyword Search logic
    # Split remaining text into paragraphs
    remaining_text = full_text[intro_size:]
    paragraphs = remaining_text.split('\n\n')
    
    keywords = [w for w in user_query.split() if len(w) > 4] # simple filter
    
    relevant_paragraphs = []
    for p in paragraphs:
        # Check if any keyword matches
        if any(k.lower() in p.lower() for k in keywords):
            relevant_paragraphs.append(p)
            
    # Add top 5 relevant paragraphs
    if relevant_paragraphs:
        context += "\n...\n[Relevant Sections Found]:\n" + "\n\n".join(relevant_paragraphs[:5])
    
    return context[:max_chars]
